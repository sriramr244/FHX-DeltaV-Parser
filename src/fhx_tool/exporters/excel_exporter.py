from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import TypeAlias, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from fhx_tool.config.fhx_schema import (
    ALARM_REPORT_COLUMN_WIDTHS,
    ALARM_REPORT_HEADERS,
    ALARM_TYPE_COLORS,
    OUTPUT_FILES,
)
from fhx_tool.domain.models import AlarmRecord, ExportedFile


ExcelCellValue: TypeAlias = str | int | float | bool | None


class ExcelExporter:
    _sheet_name = "Module Alarms"

    _header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    _header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    def export_alarm_report(
        self,
        output_directory: Path,
        records: Sequence[AlarmRecord],
    ) -> ExportedFile:
        output_directory = Path(output_directory)

        output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

        path = (
            output_directory
            / OUTPUT_FILES["alarm_report"]
        )

        workbook = Workbook()

        worksheet = cast(
            Worksheet,
            workbook.active,
        )

        worksheet.title = self._sheet_name
        worksheet.append(ALARM_REPORT_HEADERS)

        for record in records:
            worksheet.append(
                self._record_row(record)
            )

        self._format_header(worksheet)
        self._format_columns(worksheet)
        self._format_alarm_types(worksheet)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        if records:
            self._add_table(
                worksheet,
                len(records) + 1,
            )

        workbook.save(path)

        return ExportedFile(
            name=path.name,
            path=path,
            row_count=len(records),
        )

    def _record_row(
        self,
        record: AlarmRecord,
    ) -> tuple[ExcelCellValue, ...]:
        return (
            record.module_name,
            record.module_class,
            record.plant_area,
            record.controller,
            record.module_alarm,
            record.alarm_type,
            record.source_block,
            record.block_type,
            record.cause_parameter,
            record.limit_parameter,
            self._excel_value(record.limit),
            self._excel_value(record.hysteresis),
            self._excel_value(record.delay_on),
            self._excel_value(record.delay_off),
            record.enabled,
            record.priority,
            record.units,
            record.description,
        )

    def _format_header(
        self,
        worksheet: Worksheet,
    ) -> None:
        for cell in worksheet[1]:
            cell.fill = self._header_fill
            cell.font = self._header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.row_dimensions[1].height = 24

    def _format_columns(
        self,
        worksheet: Worksheet,
    ) -> None:
        for (
            column,
            width,
        ) in ALARM_REPORT_COLUMN_WIDTHS.items():
            worksheet.column_dimensions[column].width = (
                width
            )

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=False,
                )

    def _format_alarm_types(
        self,
        worksheet: Worksheet,
    ) -> None:
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            alarm_type_cell = worksheet.cell(
                row=row_index,
                column=6,
            )

            color = ALARM_TYPE_COLORS.get(
                str(alarm_type_cell.value)
            )

            if color:
                alarm_type_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=color,
                )

    def _add_table(
        self,
        worksheet: Worksheet,
        last_row: int,
    ) -> None:
        table = Table(
            displayName="ModuleAlarmTable",
            ref=f"A1:R{last_row}",
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    def _excel_value(
        self,
        value: str,
    ) -> ExcelCellValue:
        normalized = value.strip()

        if not normalized:
            return ""

        try:
            return int(normalized)
        except ValueError:
            pass

        try:
            return float(normalized)
        except ValueError:
            return normalized