"""Filtered module inventory with linked history, alarms, and decision audit."""
from collections import defaultdict
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from fhx_tool.config.fhx_schema import ALARM_REPORT_HEADERS
from fhx_tool.domain.models import AlarmRecord, ExportedFile, HistoryPoint
from fhx_tool.exporters.excel_exporter import ExcelExporter
from fhx_tool.services.module_filter import ModuleDecision


class ModuleReportExporter:
    def export(self, output_directory: Path, decisions: Sequence[ModuleDecision],
               history_points: Sequence[HistoryPoint], alarms: Sequence[AlarmRecord]) -> ExportedFile:
        output_directory.mkdir(parents=True, exist_ok=True)
        path = output_directory / "_filtered_module_report.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        history = defaultdict(list)
        for point in history_points:
            history[point.module_name].append(point)
        properties = sorted({key for point in history_points for key in point.properties})
        module_headers = ("Module", "Description", "Class", "Controller", "Plant Area", "Category", "Classification", "Reason", "History Point Count", "History Status", "History Link")
        history_headers = ("Module", "History Tag", "Module Description", "Module Class", "Plant Area", "Unit Module", "Process Cell", "Source Block", "Block Type", "History Instance", "Field", *properties)
        for status, module_sheet, history_sheet, alarm_sheet in (
            ("KEEP", "Modules", "Module History", "Module Alarms"),
            ("REVIEW", "Review Modules", "Review History", "Review Alarms"),
        ):
            selected = [d for d in decisions if d.classification == status]
            names = {d.module.module_name for d in selected}
            modules_ws = workbook.create_sheet(module_sheet)
            modules_ws.append(module_headers)
            history_ws = workbook.create_sheet(history_sheet)
            history_ws.append(history_headers)
            for decision in selected:
                module = decision.module
                points = history[module.module_name]
                first_row = history_ws.max_row + 1
                for point in points:
                    history_ws.append((point.module_name, point.history_tag, point.module_description,
                                       point.module_class, point.plant_area, point.unit_module_name,
                                       point.process_cell_name, point.source_block, point.block_type,
                                       point.history_instance, point.field_name,
                                       *(point.properties.get(key, "") for key in properties)))
                modules_ws.append((module.module_name, module.description, module.module_class,
                                   module.controller, module.plant_area, decision.category,
                                   decision.classification, decision.reason, len(points),
                                   "Configured" if points else "No history found in FHX",
                                   "View history" if points else ""))
                if points:
                    cell = modules_ws.cell(modules_ws.max_row, len(module_headers))
                    cell.hyperlink = f"#'{history_sheet}'!A{first_row}"
                    cell.style = "Hyperlink"
            alarms_ws = workbook.create_sheet(alarm_sheet)
            alarms_ws.append(ALARM_REPORT_HEADERS)
            exporter = ExcelExporter()
            for alarm in alarms:
                if alarm.module_name in names:
                    alarms_ws.append(exporter._record_row(alarm))
        audit = workbook.create_sheet("Filter Audit")
        audit.append(("Module", "Description", "Class", "Plant Area", "Classification", "Category", "Reason", "History Point Count"))
        for decision in decisions:
            module = decision.module
            audit.append((module.module_name, module.description, module.module_class,
                          module.plant_area, decision.classification, decision.category,
                          decision.reason, len(history[module.module_name])))
        for sheet in workbook:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for column in sheet.columns:
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(65, max(16, max(len(str(c.value or "")) for c in column) + 2))
                # FHX metadata is text, even when it begins with '='.
                for cell in column[1:]:
                    if cell.data_type == "f":
                        cell.data_type = "s"
        workbook.save(path)
        return ExportedFile(path.name, path, sum(d.classification == "KEEP" for d in decisions))
