from pathlib import Path

from openpyxl import load_workbook

from fhx_tool.config.fhx_schema import (
    ALARM_REPORT_HEADERS,
    OUTPUT_FILES,
)
from fhx_tool.domain.models import AlarmRecord
from fhx_tool.exporters.excel_exporter import ExcelExporter


def test_exports_formatted_alarm_workbook(
    tmp_path: Path,
) -> None:
    record = AlarmRecord(
        module_name="MODULE_001",
        module_class="CLASS_AI",
        plant_area="AREA_001/UNIT_001",
        controller="CONTROLLER_001",
        module_alarm="HI_HI_ALM",
        alarm_type="HI_HI",
        source_block="AI1",
        block_type="AI",
        cause_parameter="AI1/HI_HI_ACT",
        limit_parameter="AI1/HI_HI_LIM",
        limit="95",
        hysteresis="0.5",
        delay_on="2",
        delay_off="1",
        enabled=True,
        priority="CRITICAL",
        units="EU",
        description="Generic high-high alarm",
    )

    exported = ExcelExporter().export_alarm_report(
        tmp_path,
        [record],
    )

    expected_path = (
        tmp_path / OUTPUT_FILES["alarm_report"]
    )

    assert exported.path == expected_path
    assert exported.row_count == 1
    assert expected_path.is_file()

    workbook = load_workbook(expected_path)
    worksheet = workbook["Module Alarms"]

    headers = tuple(
        cell.value for cell in worksheet[1]
    )

    assert headers == ALARM_REPORT_HEADERS
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:R2"
    assert worksheet["A2"].value == "MODULE_001"
    assert worksheet["F2"].value == "HI_HI"
    assert worksheet["K2"].value == 95
    assert worksheet["L2"].value == 0.5
    assert worksheet["O2"].value is True
    assert "ModuleAlarmTable" in worksheet.tables


def test_exports_empty_alarm_workbook(
    tmp_path: Path,
) -> None:
    exported = ExcelExporter().export_alarm_report(
        tmp_path,
        [],
    )

    workbook = load_workbook(exported.path)
    worksheet = workbook["Module Alarms"]

    assert exported.row_count == 0
    assert worksheet.max_row == 1
    assert worksheet.freeze_panes == "A2"

    headers = tuple(
        cell.value for cell in worksheet[1]
    )

    assert headers == ALARM_REPORT_HEADERS