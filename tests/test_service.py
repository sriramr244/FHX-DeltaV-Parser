from pathlib import Path

from openpyxl import load_workbook

from fhx_tool.config.fhx_schema import OUTPUT_FILES
from fhx_tool.services.fhx_service import (
    FhxProcessingService,
)


def test_service_processes_fhx_and_exports_alarm_report(
    tmp_path: Path,
) -> None:
    source = tmp_path / "generic.fhx"

    source.write_text(
        "\n".join(
            [
                (
                    'MODULE_CLASS NAME="CLASS_AI" '
                    'CATEGORY=""'
                ),
                "{",
                '  DESCRIPTION="Generic class"',
                (
                    '  FUNCTION_BLOCK NAME="AI1" '
                    'DEFINITION="AI"'
                ),
                "  {",
                '    DESCRIPTION="Analog input block"',
                "  }",
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="AI1/HI_LIM"'
                ),
                "  {",
                "    VALUE { CV=80 }",
                "  }",
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="AI1/HI_HYS"'
                ),
                "  {",
                "    VALUE { CV=1 }",
                "  }",
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="AI1/HI_DELAY_ON"'
                ),
                "  {",
                "    VALUE { CV=2 }",
                "  }",
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="AI1/HI_DELAY_OFF"'
                ),
                "  {",
                "    VALUE { CV=0 }",
                "  }",
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="AI1/IN_SCALE"'
                ),
                "  {",
                (
                    '    VALUE { EU100=100 EU0=0 '
                    'UNITS="EU" DECPT=1 }'
                ),
                "  }",
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="HI_ALM"'
                ),
                "  {",
                "    VALUE",
                "    {",
                '      PRIORITY_NAME="WARNING"',
                "      ENAB=T",
                '      ALMATTR="AI1/HI_ACT"',
                '      LIMATTR="AI1/HI_LIM"',
                (
                    '      ALARM_DESCRIPTION='
                    '"Generic high alarm"'
                ),
                "    }",
                "  }",
                "}",
                (
                    'BATCH_EQUIPMENT_UNIT_MODULE '
                    'NAME="UNIT_001" CLASS=""'
                ),
                "{",
                "}",
                (
                    'MODULE_INSTANCE '
                    'TAG="MODULE_001" '
                    'PLANT_AREA="AREA_001/UNIT_001" '
                    'MODULE_CLASS="CLASS_AI" '
                    'CATEGORY=""'
                ),
                "{",
                '  DESCRIPTION="Generic module"',
                '  CONTROLLER="CONTROLLER_001"',
                (
                    '  ATTRIBUTE_INSTANCE '
                    'NAME="AI1/HI_LIM"'
                ),
                "  {",
                "    VALUE { CV=85 }",
                "  }",
                (
                    '  HISTORY_DATA_POINT_INSTANCE '
                    'NAME="AI1/PV"'
                ),
                "  {",
                '    HISTORY_DATA_POINT FIELD="CV"',
                "    {",
                "      ENABLED=T",
                "    }",
                "  }",
                "}",
            ]
        ),
        encoding="utf-16",
    )

    output_directory = tmp_path / "output"

    summary = FhxProcessingService().process(
        source,
        output_directory,
    )

    assert summary.history_point_count == 1
    assert summary.module_inventory_count == 1
    assert summary.alarm_count == 1
    assert summary.alarm_module_count == 1
    assert summary.alarm_report.is_file()
    filtered = load_workbook(summary.filtered_report)
    assert filtered["Modules"].max_row == 1
    assert filtered["Module History"].max_row == 1
    assert filtered["Module Alarms"].max_row == 1
    assert filtered["Review Alarms"]["A2"].value == "MODULE_001"
    assert filtered["Review History"]["A2"].value == "MODULE_001"
    history_sheet = filtered["Review History"]
    history_row = dict(zip([c.value for c in history_sheet[1]], [c.value for c in history_sheet[2]]))
    assert history_row["Source Block"] == "AI1"
    assert history_row["Block Type"] == "AI"

    history_path = (
        output_directory
        / OUTPUT_FILES["history"]
    )

    assert history_path.is_file()

    history_content = history_path.read_text(
        encoding="utf-8-sig"
    )

    assert (
        "MODULE_001/AI1/PV.CV"
        in history_content
    )

    workbook = load_workbook(
        summary.alarm_report
    )

    worksheet = workbook["Module Alarms"]

    assert worksheet["A2"].value == "MODULE_001"
    assert worksheet["B2"].value == "CLASS_AI"
    assert (
        worksheet["C2"].value
        == "AREA_001/UNIT_001"
    )
    assert (
        worksheet["D2"].value
        == "CONTROLLER_001"
    )
    assert worksheet["E2"].value == "HI_ALM"
    assert worksheet["F2"].value == "HI"
    assert worksheet["G2"].value == "AI1"
    assert worksheet["H2"].value == "AI"
    assert (
        worksheet["I2"].value
        == "AI1/HI_ACT"
    )
    assert (
        worksheet["J2"].value
        == "AI1/HI_LIM"
    )
    assert worksheet["K2"].value == 85
    assert worksheet["L2"].value == 1
    assert worksheet["M2"].value == 2
    assert worksheet["N2"].value == 0
    assert worksheet["O2"].value is True
    assert worksheet["P2"].value == "WARNING"
    assert worksheet["Q2"].value == "EU"
    assert (
        worksheet["R2"].value
        == "Generic high alarm"
    )

def test_filtered_report_associates_history_and_preserves_review(tmp_path):
    source = tmp_path / "filtered.fhx"
    source.write_text('''MODULE_CLASS NAME="AI" CATEGORY=""
{
 HISTORY_DATA_POINT_INSTANCE NAME="AI1/PV"
 {
  HISTORY_DATA_POINT FIELD="CV"
  {
   ENABLED=T
   SAMPLE_PERIOD_SECONDS=2
  }
 }
}
MODULE_INSTANCE TAG="PT101" MODULE_CLASS="AI" PLANT_AREA="AREA" CATEGORY=""
{
 DESCRIPTION="Pressure"
}
MODULE_INSTANCE TAG="SS101" MODULE_CLASS="AI" PLANT_AREA="AREA" CATEGORY=""
{
 DESCRIPTION="Signal select"
}
MODULE_INSTANCE TAG="MYSTERY" MODULE_CLASS="AI" PLANT_AREA="AREA" CATEGORY=""
{
 DESCRIPTION="Unknown"
}
MODULE TAG="TT101" PLANT_AREA="AREA" CATEGORY=""
{
 DESCRIPTION="Temperature"
}
''', encoding="utf-16")
    summary = FhxProcessingService().process(source, tmp_path / "output")
    assert (summary.kept_module_count, summary.review_module_count, summary.removed_module_count) == (2, 1, 1)
    assert summary.filtered_report in summary.exported_files
    workbook = load_workbook(summary.filtered_report)
    modules = workbook["Modules"]
    assert [row[0] for row in modules.iter_rows(min_row=2, values_only=True)] == ["PT101", "TT101"]
    assert modules["I2"].value == 1
    assert modules["K2"].hyperlink.target == "#'Module History'!A2"
    assert modules["I3"].value == 0
    assert modules["J3"].value == "No history found in FHX"
    history = workbook["Module History"]
    assert history.max_row == 2
    assert history["B2"].value == "PT101/AI1/PV.CV"
    assert dict(zip([c.value for c in history[1]], [c.value for c in history[2]]))["SAMPLE_PERIOD_SECONDS"] == "2"
    headers = [c.value for c in history[1]]
    row = dict(zip(headers, [c.value for c in history[2]]))
    assert row["Source Block"] == "AI1"
    assert row["Block Type"] is None  # A history path alone cannot prove an AI block.
    assert row["History Instance"] == "AI1/PV"
    assert row["Field"] == "CV"
    assert workbook["Review History"]["A2"].value == "MYSTERY"
    assert workbook["Filter Audit"].max_row == 5
